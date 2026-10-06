"""Build the editable paper with tables mechanically populated from evidence.

Run with the Codex bundled Python document runtime. No campaign is modified.
"""
from pathlib import Path
import json
import re
import hashlib
import csv
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'paper'
load=lambda f:json.loads((ROOT/f).read_text(encoding='utf-8'))
selection=load('results/ml-v7-models-complete/selection.json');summary=load('results/ml-v7-evaluation/summary.json');contrasts=load('results/ml-v7-evaluation/contrasts.json')
mods=list(selection['models']);names={'RSSI':'RSSI','AoA2':'AoA2','AoA4':'AoA4','PBR':'PBR','RTT':'RTT','PBR_RTT':'PBR/RTT','UWB':'UWB'}
family_names={'logistic':'Logistic','random_forest':'Forest','boosted_trees':'Boosted','gaussian_bayes':'Gaussian Bayes'}
def row(split,mod,det,condition='all'):
    return next(r for r in summary if (r['split'],r['modality'],r['detector'],r['condition'])==(split,mod,det,condition))
def fmt(n):return '—' if n is None else f'{n:.3f}'
tables={}
def add(key,caption,headers,body,widths=None):tables[key]={'caption':caption,'headers':headers,'rows':body,'widths':widths}
add('budgets','Table 1 Observation configurations', ['Budget','Additional measurements','Receiver configuration'],[
['RSSI','None','Four RSSI gateways'],['AoA2','Bearings','g0 and g2 upgraded; RSSI at all four'],['AoA4','Bearings','All four gateways upgraded'],
['PBR','Multi-tone reciprocal phase','Four ranging anchors'],['RTT','Elapsed and turnaround time','Four ranging anchors'],['PBR/RTT','PBR plus RTT consistency gate','Four ranging anchors'],['UWB','DS-TWR intervals','Four ranging anchors']], [1.0,2.4,3.3])
add('channel','Table 2 Assumed sensor and reception parameters', ['Parameter','Ideal setting','Wall setting'],[
['RSSI noise SD (dB)','6','6'],['RSSI and angular AR correlation','0','0.8'],['Persistent device RSSI bias SD (dB)','0','2'],['Wall attenuation (dB)','0','6'],
['Bearing noise SD (degrees)','5','5'],['Gateway orientation bias SD (degrees)','0','2'],['Wall bearing bias (degrees)','0','8'],
['Independent packet loss','0.10','0.10'],['Sensitivity floor (dBm)','−95','−95'],['Burst-start probability per link per second','0','0.008'],['Burst length (s)','0','3'],
['PBR phase noise SD (rad)','0.25','0.50'],['Blocked-link reflected amplitude / excess path','0 / 0','1.4 / 2 m'],
['RTT elapsed-time noise SD (ns)','10','10'],['UWB interval noise SD (ns)','0.4','0.4'],['RTT blocked-link bias (m)','0','1.5'],['UWB blocked-link bias (m)','0','0.8'],
['Persistent PBR / RTT / UWB range-bias SD (m)','0.15 / 0.30 / 0.10','0.15 / 0.30 / 0.10']], [4.0,1.35,1.35])
add('models','Table 3 Declared classifier search grids',['Family','Candidate settings','Fixed fitting settings'],[
['Logistic','C ∈ {0.1, 1}','Standardized; balanced weights; max 2,000 iterations'],['Random forest','Depth ∈ {8, 12}','64 trees; leaf minimum 25; balanced weights'],
['Histogram boosting','Leaves ∈ {15, 31}','80 iterations; leaf minimum 50; L2 = 0.1; balanced weights; no early stopping'],
['Gaussian Bayes','Variance smoothing ∈ {10⁻⁶, 10⁻³}','Standardized; empirical class priors']], [1.35,1.9,3.45])
partition_rows=[]
for s,label in [('train','Training'),('calibration','Calibration'),('validation','Validation'),('test','Primary test'),('external','Unseen motion')]:
    m=load(f'results/ml-v7-data/{s}/manifest.json');unique={c['seed']:c for c in m['cases']}
    partition_rows.append([label,f"{m['seeds'][0]}–{m['seeds'][-1]}",str(len(unique)),str(sum(c['referenceContacts'] for c in unique.values())),f"{sum(c['positiveFrames'] for c in unique.values()):,}"])
add('partitions','Table 4 Independent motion support by partition',['Partition','Seed range','Recordings','Unique contacts','Positive pair-frames'],partition_rows,[1.4,1.35,1.0,1.25,1.7])
primary=[];family_rows=[];conditions=[];abla=[];quality=[];chosen=[]
for mod in mods:
    det=selection['best_by_validation'][mod];a=row('test',mod,det);r=row('test',mod,'tuned_rule');o=row('external',mod,det)
    c=next(c for c in contrasts if (c['split'],c['modality'],c['contrast'])==('test',mod,f'{det} minus tuned_rule'))
    ci=c['pairedBootstrap95']
    primary.append([names[mod],family_names[det],fmt(r['meanSeed_f1']),fmt(a['meanSeed_f1']),f"{a['tp']}/{a['fp']}/{a['fn']}",f"{c['meanSeedF1Difference']:.3f}\n[{ci[0]:.3f}, {ci[1]:.3f}]"])
    family_rows.append([names[mod]]+[fmt(row('test',mod,d)['meanSeed_f1']) for d in ['strict_rule','temporal_rule','tuned_rule','logistic','random_forest','boosted_trees','gaussian_bayes']])
    conditions.append([names[mod],fmt(row('test',mod,det,'LOS')['meanSeed_f1']),fmt(row('test',mod,det,'wall-multipath')['meanSeed_f1']),fmt(o['meanSeed_f1']),f"{o['tp']}/{o['fp']}/{o['fn']}",fmt(o['meanSeed_falseSessionsPerPairHour'])])
    abla.append([names[mod],fmt(a['meanSeed_f1'])]+[fmt(row('test',mod,d)['meanSeed_f1']) for d in ['learned_markov','without_contact_history','without_room_feature']])
    quality.append([names[mod],fmt(a['meanSeed_frameBrier']),fmt(a['meanSeed_frameECE']),fmt(a['meanSeed_missedContactTimeRate']),fmt(a['meanSeed_falseSessionsPerPairHour']),f"{a['meanSeed_durationMAE']:.1f}",f"{a['meanSeed_medianDetectionLatency']:.1f}"])
    settings=selection['models'][mod][det];rule=selection['tuned_rules'][mod]
    chosen.append([names[mod],family_names[det],settings['config'],f"{settings['threshold']} / {settings['decoder']}",f"{rule['distance']} m / {rule['decoder']}"])
add('primary','Table 5 Validation-selected primary contact reconstruction',['Budget','Family','Rule F1','ML F1','TP/FP/FN','ML minus rule\n95% interval'],primary,[.7,.85,.7,.7,1.2,2.05])
add('families','Table 6 Primary contact F1 across full families and controls',['Budget','Strict','Temporal','Tuned','Logistic','Forest','Boosted','Bayes'],family_rows,[.85,.8,.85,.85,.85,.85,.85,.85])
add('conditions','Table 7 Selected-model conditions and unseen motion',['Budget','Ideal F1','Wall F1','Unseen F1','Unseen\nTP/FP/FN','Unseen FP\nper pair-hour'],conditions,[.85,.85,.85,.95,1.5,1.7])
add('ablations','Table 8 Primary contact F1 for temporal extension and ablations',['Budget','Selected full','Markov','No history','No room feature'],abla,[1,1.4,1.2,1.4,1.7])
add('quality','Table 9 Selected-model probability and time outcomes',['Budget','Brier','ECE','Missed\ntime rate','FP per\npair-hour','Duration\nMAE s','Confirm\ndelay s'],quality,[.75,.75,.75,1.05,1.1,1.15,1.15])
service=[]
for line in (ROOT/'docs/RANGING_V6_RESULTS.md').read_text(encoding='utf-8').splitlines():
    cells=[s.strip() for s in line.strip('|').split('|')]
    if len(cells)==10 and cells[2]=='uwb-ds-twr':service.append([cells[0],cells[1],cells[6],cells[9]])
assert len(service)==6
add('service','Table 10 Separate UWB service-load study',['Channel setting','Scheduled clients','Mean contact F1','Availability'],service,[2.3,1.5,1.45,1.45])
descriptions=[
('Estimated separation capped at 30; zero if unavailable','m'),('Same side of public wall from estimated positions','binary'),('Tracked pair positions and BLE ages at most 3 s','binary'),
('Smaller latest-telemetry age; 30 if none','s'),('Larger latest-telemetry age; 30 if none','s'),('Common receivers in five-second telemetry histories','count'),
('Smaller device mean of receiver-median RSSI','dBm'),('Larger device mean of receiver-median RSSI','dBm'),('Absolute mean signed RSSI difference on common receivers','dB'),
('Mean absolute common-receiver RSSI differences after offset removal','dB'),('Maximum offset-removed RSSI profile difference','dB'),
('Smaller received-observation count over five-second histories','count'),('Larger received-observation count over five-second histories','count'),
('Mean circular angular difference on common enabled receivers','rad'),('Maximum circular angular difference','rad'),
('Mean absolute common-anchor range difference','m'),('Maximum absolute common-anchor range difference','m'),('Mean absolute observed-versus-estimated range residual','m'),
('Mean unwrapped phase-slope residual on current PBR links','rad'),('Mean PBR/RTT disagreement on accepted combined links','m'),
('Mean available separation in preceding ten seconds','m'),('SD of available separation in preceding ten seconds','m'),('Mean available separation in preceding thirty seconds','m'),
('Fixed estimated 2 m same-room positives / history samples over 10 s','fraction'),('Same fraction over 30 s','fraction'),
('Available samples / history samples over 10 s','fraction'),('Same fraction over 30 s','fraction'),('Last minus first available distance in ten-second history','m'),
('Mean offset-removed RSSI profile difference in ten-second history','dB')]
assert len(descriptions)==len(selection['features'])==29
feature_rows=[[str(i+1),n,d,'0–1' if u=='fraction' else u] for i,(n,(d,u)) in enumerate(zip(selection['features'],descriptions))]
add('features','Table A1 Observable feature schema',['No','Column name','Definition','Unit'],feature_rows,[.35,2.05,3.85,.45])
add('selections','Table C1 Frozen primary model and rule settings',['Budget','Family','Hyperparameter','ML threshold / decoder','Rule distance / decoder'],chosen,[.75,.9,1.35,1.85,1.85])

refs=[
('P. Vanhems et al.','Estimating Potential Infection Transmission Routes in Hospital Wards Using Wearable Proximity Sensors','PLoS ONE, 8(9), e73970, 2013','https://doi.org/10.1371/journal.pone.0073970'),
('D. J. Leith and S. Farrell','Coronavirus Contact Tracing: Evaluating the Potential of Using Bluetooth Received Signal Strength for Proximity Detection','arXiv:2006.06822, 2020','https://arxiv.org/abs/2006.06822'),
('Bluetooth SIG','Direction Finding','Technical resource, accessed 6 October 2026','https://www.bluetooth.com/learn-about-bluetooth/feature-enhancements/direction-finding/'),
('Bluetooth SIG','Bluetooth Channel Sounding','Technical resource, accessed 6 October 2026','https://www.bluetooth.com/learn-about-bluetooth/feature-enhancements/channel-sounding/'),
('Bluetooth SIG','Bluetooth Channel Sounding A Technical Overview','Technical resource, accessed 6 October 2026','https://www.bluetooth.com/channel-sounding-tech-overview/'),
('P. Rathje and O. Landsiedel','Precise Ranging: Modeling Bias and Variance of Double-Sided Two-Way Ranging with TDoA Extraction under Multipath and NLOS Effects','arXiv:2410.12826v2, 2024','https://arxiv.org/abs/2410.12826v2'),
('A. Niculescu-Mizil and R. Caruana','Predicting Good Probabilities with Supervised Learning','Proceedings of ICML, pp. 625–632, 2005; DOI 10.1145/1102351.1102430','https://icml.cc/Conferences/2005/proceedings/papers/079_GoodProbabilities_NiculescuMizilCaruana.pdf'),
('C. Guo, G. Pleiss, Y. Sun, and K. Q. Weinberger','On Calibration of Modern Neural Networks','Proceedings of ICML, PMLR 70, pp. 1321–1330, 2017','https://proceedings.mlr.press/v70/guo17a.html'),
('F. Pedregosa et al.','Scikit-learn: Machine Learning in Python','Journal of Machine Learning Research, 12, pp. 2825–2830, 2011','https://www.jmlr.org/papers/v12/pedregosa11a.html'),
('Scikit-learn developers','Cross-validation Evaluating Estimator Performance','Software documentation, accessed 6 October 2026','https://scikit-learn.org/1.7/modules/cross_validation.html'),
('Scikit-learn developers','HistGradientBoostingClassifier','Version 1.7.2 documentation, accessed 6 October 2026','https://scikit-learn.org/1.7/modules/generated/sklearn.ensemble.HistGradientBoostingClassifier.html'),
('Qorvo','DWM3001C','Product technical information, accessed 6 October 2026','https://www.qorvo.com/products/p/DWM3001C'),
('P. J. Huber','Robust Estimation of a Location Parameter','The Annals of Mathematical Statistics, 35(1), pp. 73–101, 1964','https://doi.org/10.1214/aoms/1177703732'),
('R. E. Kalman','A New Approach to Linear Filtering and Prediction Problems','Journal of Basic Engineering, 82(1), pp. 35–45, 1960','https://doi.org/10.1115/1.3662552'),
('L. Breiman','Random Forests','Machine Learning, 45, pp. 5–32, 2001','https://doi.org/10.1023/A:1010933404324'),
('J. H. Friedman','Greedy Function Approximation: A Gradient Boosting Machine','The Annals of Statistics, 29(5), pp. 1189–1232, 2001','https://doi.org/10.1214/aos/1013203451')]
reproduction='''python -m venv .research-venv
.\\.research-venv\\Scripts\\python.exe -m pip install -r research/ml-v7/requirements.txt --index-url https://pypi.org/simple
node --test research/ml-v7/test.mjs
.\\.research-venv\\Scripts\\python.exe research/ml-v7/tests.py
node research/ml-v7/generate.mjs train results/replay-data/train
node research/ml-v7/generate.mjs calibration results/replay-data/calibration
node research/ml-v7/generate.mjs validation results/replay-data/validation
.\\.research-venv\\Scripts\\python.exe research/ml-v7/train.py results/replay-data results/replay-models
node research/ml-v7/generate.mjs test results/replay-data/test
node research/ml-v7/generate.mjs external results/replay-data/external
.\\.research-venv\\Scripts\\python.exe research/ml-v7/evaluate.py results/replay-data results/replay-models results/replay-evaluation
.\\.research-venv\\Scripts\\python.exe research/ml-v7/verify.py results/replay-data results/replay-models results/replay-evaluation
.\\.research-venv\\Scripts\\python.exe research/ml-v7/report.py results/replay-models results/replay-evaluation results/replay-report --data results/replay-data'''

def reference_text(i,author,title,venue,url):
    lead=f'[{i}] {author.rstrip(".")}, “{title},” '
    bibliographies={
        1:'PLoS ONE, vol. 8, no. 9, Art. no. e73970, 2013, doi: 10.1371/journal.pone.0073970.',
        2:'arXiv:2006.06822, 2020.',
        6:'arXiv:2410.12826v2, 2024.',
        7:'in Proc. International Conference on Machine Learning (ICML), 2005, pp. 625–632, doi: 10.1145/1102351.1102430.',
        8:'in Proc. International Conference on Machine Learning (ICML), Proc. Machine Learning Research, vol. 70, 2017, pp. 1321–1330.',
        9:'Journal of Machine Learning Research, vol. 12, 2011, pp. 2825–2830.',
        13:'The Annals of Mathematical Statistics, vol. 35, no. 1, 1964, pp. 73–101, doi: 10.1214/aoms/1177703732.',
        14:'Journal of Basic Engineering, vol. 82, no. 1, 1960, pp. 35–45, doi: 10.1115/1.3662552.',
        15:'Machine Learning, vol. 45, 2001, pp. 5–32, doi: 10.1023/A:1010933404324.',
        16:'The Annals of Statistics, vol. 29, no. 5, 2001, pp. 1189–1232, doi: 10.1214/aos/1013203451.'}
    if i in bibliographies:
        body=bibliographies[i]
        if i in [2,6,8,9]:body+=f' [Online]. Available: {url}'
    else:
        body=('Version 1.7.2 software documentation. ' if i==11 else 'Technical documentation. ')
        body+=f'[Online]. Available: {url}. Accessed: Oct. 6, 2026.'
    return lead+body

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.28)
for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    style=doc.styles[style_name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    for old in list(rpr.findall(qn('w:rFonts'))):rpr.remove(old)
    fonts=OxmlElement('w:rFonts')
    for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
    rpr.insert(0,fonts)
    ppr=style.element.get_or_add_pPr()
    for old in list(ppr.findall(qn('w:pBdr'))):ppr.remove(old)
    style.font.size=Pt(11)
    style.paragraph_format.space_after=Pt(5);style.paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(18);doc.styles['Title'].font.bold=True;doc.styles['Title'].paragraph_format.space_after=Pt(8)
for name,size in [('Heading 1',12.5),('Heading 2',11.5)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True;doc.styles[name].paragraph_format.space_before=Pt(11);doc.styles[name].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(9.5);doc.styles['Caption'].font.italic=True;doc.styles['Caption'].font.bold=False
normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
doc.core_properties.title='Learning Sustained Contacts from RSSI, Angular Measurements, and Ranging: A Reproducible Simulation Study'
doc.core_properties.author='Arjun K. V.; Nabil Ahmad Khan; Md. Aasim Shaibani; Ayyan Anwar; Alakananda K.'
doc.core_properties.subject='Synthetic contact reconstruction with grouped validation and calibrated classifiers'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
def text(p,value):
    run=p.add_run(value);run.font.name='Times New Roman'
    return run
def table(record):
    caption=doc.add_paragraph(record['caption'],style='Caption');caption.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(record['headers']));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=record['widths'] or [6.7/len(record['headers'])]*len(record['headers'])
    scale=6.7/sum(widths);widths=[w*scale for w in widths]
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for cell,w in zip(t.rows[0].cells,widths):cell.width=Inches(w)
    trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    for values in [record['headers'],*record['rows']]:
        cells=t.rows[0].cells if values is record['headers'] else t.add_row().cells
        header=values is record['headers']
        for j,(cell,value,w) in enumerate(zip(cells,values,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                item=OxmlElement(f'w:{side}');item.set(qn('w:w'),'65');item.set(qn('w:type'),'dxa');margins.append(item)
            tcpr.append(margins);borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement(f'w:{edge}');el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            if header:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6E6E6');tcpr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.0
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(str(value))>25 else WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(str(value));r.font.name='Times New Roman';r.font.size=Pt(9 if len(record['headers'])<=6 else 8.5);r.bold=header
            p.paragraph_format.keep_with_next=False
        no_split=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(no_split)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
def equation(expression,number):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7)
    math=OxmlElement('m:oMath');mr=OxmlElement('m:r');prop=OxmlElement('m:rPr');normal=OxmlElement('m:nor');prop.append(normal);mr.append(prop)
    wrpr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');wrpr.append(fonts)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'20');wrpr.append(size);mr.append(wrpr)
    content=OxmlElement('m:t');content.text=expression;mr.append(content);math.append(mr);p._p.append(math)
    r=p.add_run(f'   ({number})');r.font.size=Pt(10)
def figure(file,caption):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(OUT/'figures'/file),width=Inches(4.6 if 'floorplan' in file else 6.65))
    cp=doc.add_paragraph(caption,style='Caption');cp.paragraph_format.keep_with_next=False
def code():
    # Wrap lengthy Windows commands to prevent right-margin overflow; each
    # line remains executable with PowerShell's explicit continuation marker.
    for line in reproduction.splitlines():
        if len(line)>88 and ' ' in line:
            words=line.split(' ');parts=[];current=''
            for word in words:
                if len(current)+len(word)>85:parts.append(current+' `');current='    '+word
                else:current+=((' ' if current else '')+word)
            parts.append(current)
        else:parts=[line]
        for part in parts:
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
            r=p.add_run(part);r.font.name='Consolas';r.font.size=Pt(8.5)

source=(OUT/'manuscript_content.md').read_text(encoding='utf-8');expanded=[]
for block in re.split(r'\n\s*\n',source.strip()):
    block=block.strip()
    if block.startswith('@@TABLE:'):
        key=block.removeprefix('@@TABLE:').removesuffix('@@');record=tables[key];table(record)
        expanded.append(record['caption']+'\n\n'+'| '+' | '.join(record['headers']).replace('\n',' ')+' |\n| '+' | '.join(['---']*len(record['headers']))+' |\n'+'\n'.join('| '+' | '.join(map(str,r)).replace('\n',' ')+' |' for r in record['rows']))
    elif block.startswith('@@FIGURE:'):
        file,caption=block.removeprefix('@@FIGURE:').removesuffix('@@').split('|',1);figure(file,caption);expanded.append(f'![{caption}](figures/{file})')
    elif block.startswith('@@EQUATION:'):
        expression,number=block.removeprefix('@@EQUATION:').removesuffix('@@').rsplit('|',1);equation(expression,number);expanded.append(f'{expression}   ({number})')
    elif block=='@@REFERENCES@@':
        for i,(author,title,venue,url) in enumerate(refs,1):
            p=doc.add_paragraph(reference_text(i,author,title,venue,url));p.paragraph_format.space_after=Pt(6)
            for r in p.runs:r.font.size=Pt(9.5)
        expanded.append('\n\n'.join(reference_text(i,a,t,v,u) for i,(a,t,v,u) in enumerate(refs,1)))
    elif block.startswith('@@CODE:'):code();expanded.append('```powershell\n'+reproduction+'\n```')
    elif block.startswith('# '):doc.add_paragraph(block[2:],'Title');expanded.append(block)
    elif block.startswith('### '):doc.add_paragraph(block[4:],'Heading 2');expanded.append(block)
    elif block.startswith('## '):
        h=doc.add_paragraph(block[3:],'Heading 1')
        if block in ['## References','## Appendix A Feature schema and implementation conventions']:h.paragraph_format.page_break_before=True
        expanded.append(block)
    else:
        p=doc.add_paragraph(block.replace('\n',' '));p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if block.startswith('Arjun K.') or block.startswith('Department '):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.font.size=Pt(10.5)
        if block.startswith('Keywords:'):
            for r in p.runs:r.font.size=Pt(10)
        expanded.append(block)
final=OUT/'Contact_Reconstruction_Research_Paper.docx';doc.save(final)
(OUT/'Contact_Reconstruction_Research_Paper.md').write_text('\n\n'.join(expanded)+'\n',encoding='utf-8')
(OUT/'manuscript_tables.json').write_text(json.dumps(tables,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
references=[{'number':i,'authors':a,'title':t,'venue':v,'url':u,'formatted':reference_text(i,a,t,v,u)} for i,(a,t,v,u) in enumerate(refs,1)]
(OUT/'references.json').write_text(json.dumps(references,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
assert '@@' not in (OUT/'Contact_Reconstruction_Research_Paper.md').read_text(encoding='utf-8')
print(f'Created detailed manuscript: {final}; {len(tables)} tables, 4 figures, 11 equations, {len(refs)} references.')
